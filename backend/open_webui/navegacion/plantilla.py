"""Construcción y cache de la plantilla Jinja (docxtpl) de la Declaración Responsable.

Copia el DOCX editable oficial del docs_pack e inserta tags Jinja2. Nunca
modifica el original.
"""

from __future__ import annotations

import copy
import hashlib
import logging
import os
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from open_webui.env import DATA_DIR
from open_webui.navegacion.geozona import ruta_docs_pack

log = logging.getLogger(__name__)

DEFAULT_EDITABLE_RELATIVE = (
    Path('Formularios, Instrucciones de cumplimentación y anexos')
    / 'EDITABLE'
    / '1. FORMULARIO Declaración responsable Miño-Sil.docx'
)
DEFAULT_PLANTILLA_PATH = DATA_DIR / 'plantillas' / 'plantilla_declaracion_responsable.docx'
METADATA_VERSION = 1

FUENTE_CASILLAS = 'Wingdings'
MARCA_SI = 'þ'  # Wingdings 0xFE
MARCA_NO = 'o'  # Wingdings 0x6F


def ruta_editable() -> Path:
    override = os.getenv('DR_EDITABLE_DOCX_PATH', '').strip()
    if override:
        return Path(override)
    return ruta_docs_pack() / DEFAULT_EDITABLE_RELATIVE


def ruta_plantilla() -> Path:
    override = os.getenv('DR_PLANTILLA_PATH', '').strip()
    if override:
        return Path(override)
    return Path(DEFAULT_PLANTILLA_PATH)


def ruta_metadata_plantilla(plantilla: Path | None = None) -> Path:
    return (plantilla or ruta_plantilla()).parent / 'plantilla_metadata.json'


def plantilla_disponible(plantilla: Path | None = None) -> bool:
    path = plantilla or ruta_plantilla()
    return path.is_file() and path.stat().st_size > 0


def _huella_archivo(path: Path) -> dict[str, Any]:
    digest = hashlib.sha256()
    with path.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b''):
            digest.update(chunk)
    st = path.stat()
    return {
        'version': METADATA_VERSION,
        'source': str(path.resolve()),
        'sha256': digest.hexdigest(),
        'size': st.st_size,
        'mtime_ns': st.st_mtime_ns,
    }


def _cargar_metadata(plantilla: Path | None = None) -> dict | None:
    meta_path = ruta_metadata_plantilla(plantilla)
    if not meta_path.is_file():
        return None
    try:
        import json

        return json.loads(meta_path.read_text(encoding='utf-8'))
    except Exception:
        return None


def _guardar_metadata(huella: dict, plantilla: Path | None = None) -> None:
    import json

    meta_path = ruta_metadata_plantilla(plantilla)
    meta_path.parent.mkdir(parents=True, exist_ok=True)
    meta_path.write_text(json.dumps(huella, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def necesita_regenerar_plantilla(
    editable: Path | None = None,
    plantilla: Path | None = None,
) -> bool:
    src = editable or ruta_editable()
    dst = plantilla or ruta_plantilla()
    if not dst.is_file():
        return True
    if not src.is_file():
        return False
    meta = _cargar_metadata(dst)
    if not meta or meta.get('version') != METADATA_VERSION:
        return True
    actual = _huella_archivo(src)
    return (
        meta.get('sha256') != actual['sha256']
        or meta.get('size') != actual['size']
        or meta.get('source') != actual['source']
    )


def expr_casilla(condicion: str) -> str:
    return f"{{{{ '{MARCA_SI}' if {condicion} else '{MARCA_NO}' }}}}"


def append_text(paragraph, text):
    ref_run = paragraph.runs[0] if paragraph.runs else None
    run = paragraph.add_run(text)
    if ref_run is not None:
        run.font.name = ref_run.font.name
        run.font.size = ref_run.font.size
    return run


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    ref_run = paragraph.runs[0] if paragraph.runs else None
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    if ref_run is not None:
        run.font.name = ref_run.font.name
        run.font.size = ref_run.font.size
    return run


def replace_paragraph_text(paragraph, text):
    ref_run = paragraph.runs[0] if paragraph.runs else None
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    if ref_run is not None:
        run.font.name = ref_run.font.name
        run.font.size = ref_run.font.size
    return run


def forzar_fuente(run_element, nombre):
    rfonts = run_element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(atributo), nombre)
    return run_element


def _run_has_sym_child(run):
    return run._element.find(qn('w:sym')) is not None


def process_checkbox_group(paragraph, expr_builder, sufijo_builder=None):
    runs = list(paragraph.runs)
    groups = []
    current = None
    for r in runs:
        if r.text == '' and _run_has_sym_child(r):
            current = {'sym': r, 'label_runs': []}
            groups.append(current)
        elif current is not None:
            current['label_runs'].append(r)

    labels = []
    for group in groups:
        label_runs = group['label_runs']
        if not label_runs:
            continue
        raw_text = ''.join(r.text for r in label_runs)
        clean_label = raw_text.strip()
        if not clean_label:
            continue
        labels.append(clean_label)
        sym_run = group['sym']
        sym_run.text = expr_builder(clean_label)
        forzar_fuente(sym_run._element, FUENTE_CASILLAS)
        if sufijo_builder is not None:
            sufijo = sufijo_builder(clean_label)
            if sufijo:
                label_runs[-1].text = label_runs[-1].text + sufijo
    return labels


def replace_sdt_with_text(sdt_element, text, fuente=FUENTE_CASILLAS):
    sdt_pr = sdt_element.find(qn('w:sdtPr'))
    rpr_source = sdt_pr.find(qn('w:rPr')) if sdt_pr is not None else None
    new_run = OxmlElement('w:r')
    if rpr_source is not None:
        new_run.append(copy.deepcopy(rpr_source))
    new_t = OxmlElement('w:t')
    new_t.set(qn('xml:space'), 'preserve')
    new_t.text = text
    new_run.append(new_t)
    if fuente:
        forzar_fuente(new_run, fuente)
    sdt_element.addnext(new_run)
    parent = sdt_element.getparent()
    if parent is not None:
        parent.remove(sdt_element)
    return new_run


def replace_leading_sym_with_mark(paragraph, mark):
    for run in paragraph.runs:
        if run.text == '' and _run_has_sym_child(run):
            run.text = mark
            forzar_fuente(run._element, FUENTE_CASILLAS)
            return True
    return False


def condicionar_anio(paragraph, texto_hueco, condicion, variable='plazo_anio_yy'):
    runs = list(paragraph.runs)
    indices_hueco = [i for i, r in enumerate(runs) if r.text == texto_hueco]
    if len(indices_hueco) != 1:
        raise RuntimeError(
            f'Se esperaba exactamente 1 run con texto {texto_hueco!r} en '
            f'{paragraph.text!r}, se encontraron {len(indices_hueco)}.'
        )
    hueco = runs[indices_hueco[0]]
    anteriores = reversed(runs[: indices_hueco[0]])
    previo = next((r for r in anteriores if r.text.strip()), None)
    if previo is None or not previo.text.rstrip().endswith('20'):
        raise RuntimeError(
            "No se encontró el '20' impreso justo antes del hueco del año en: "
            f'{paragraph.text!r}'
        )
    sin_cola = previo.text.rstrip()
    previo.text = sin_cola[:-2] + previo.text[len(sin_cola) :]
    hueco.text = '{% if ' + condicion + ' %}20{{ ' + variable + ' }}{% endif %}'


def replace_textbox_matching(cell, needle, new_text):
    reemplazados = 0
    for paragraph in cell._tc.findall('.//' + qn('w:txbxContent') + '//' + qn('w:p')):
        nodos_t = paragraph.findall('.//' + qn('w:t'))
        joined = ''.join(nodo.text or '' for nodo in nodos_t)
        if needle not in joined or not nodos_t:
            continue
        nodos_t[0].text = new_text
        for nodo in nodos_t[1:]:
            nodo.text = ''
        reemplazados += 1
    return reemplazados


def build_plantilla(
    editable: Path | None = None,
    plantilla: Path | None = None,
) -> dict[str, Any]:
    """Copia el DOCX editable e inserta tags Jinja. Devuelve resumen de tags."""
    src = Path(editable) if editable is not None else ruta_editable()
    dst = Path(plantilla) if plantilla is not None else ruta_plantilla()

    if not src.is_file():
        raise FileNotFoundError(
            f'No se encontró el formulario editable en {src}. '
            'Espera al seed de docs_pack o configura DR_EDITABLE_DOCX_PATH.'
        )

    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(src, dst)

    doc = Document(str(dst))
    tables = doc.tables
    embalses_taggeados: list[str] = []

    # Sección 1 — Declarante (tabla 0)
    t0 = tables[0]
    header_p = t0.rows[0].cells[0].paragraphs[0]
    sdts = header_p._p.findall('.//' + qn('w:sdt'))
    declarante_tipo_opciones = ['PERSONA FÍSICA', 'PERSONA JURÍDICA', 'ENTIDAD PÚBLICA']
    if len(sdts) != 3:
        raise RuntimeError(
            f'Se esperaban 3 casillas w:sdt para declarante_tipo, se encontraron {len(sdts)}'
        )
    for sdt_el, opcion in zip(sdts, declarante_tipo_opciones):
        replace_sdt_with_text(sdt_el, expr_casilla(f"declarante_tipo == '{opcion}'") + ' ')

    append_text(t0.rows[1].cells[0].paragraphs[0], ' {{ declarante_nombre }}')
    append_text(t0.rows[1].cells[5].paragraphs[0], ' {{ declarante_dni }}')
    append_text(t0.rows[2].cells[0].paragraphs[0], ' {{ declarante_direccion }}')
    append_text(t0.rows[2].cells[2].paragraphs[0], ' {{ declarante_numero }}')
    append_text(t0.rows[2].cells[3].paragraphs[0], ' {{ declarante_bloque }}')
    append_text(t0.rows[2].cells[4].paragraphs[0], ' {{ declarante_escalera }}')
    append_text(t0.rows[2].cells[6].paragraphs[0], ' {{ declarante_planta }}')
    append_text(t0.rows[2].cells[7].paragraphs[0], ' {{ declarante_puerta }}')
    append_text(t0.rows[2].cells[8].paragraphs[0], ' {{ declarante_cp }}')
    append_text(t0.rows[3].cells[0].paragraphs[0], ' {{ declarante_provincia }}')
    append_text(t0.rows[3].cells[1].paragraphs[0], ' {{ declarante_municipio }}')
    append_text(t0.rows[3].cells[2].paragraphs[0], ' {{ declarante_localidad }}')
    append_text(t0.rows[3].cells[7].paragraphs[0], ' {{ declarante_pais }}')
    append_text(t0.rows[4].cells[0].paragraphs[0], ' {{ declarante_email }}')
    append_text(t0.rows[4].cells[5].paragraphs[0], ' {{ declarante_telefono }}')

    # Sección 3 — Notificación (tabla 2)
    t2 = tables[2]
    dest_declarante_p = t2.rows[1].cells[0].paragraphs[0]
    dest_repr_p = t2.rows[1].cells[0].paragraphs[1]
    medio_elec_p = t2.rows[1].cells[1].paragraphs[0]
    medio_papel_p = t2.rows[1].cells[1].paragraphs[1]

    def _replace_unique_sdt(paragraph, mark):
        sdts_p = paragraph._p.findall('.//' + qn('w:sdt'))
        if len(sdts_p) != 1:
            raise RuntimeError(
                f'Se esperaba 1 casilla w:sdt en notificación, se encontraron {len(sdts_p)}'
            )
        replace_sdt_with_text(sdts_p[0], mark)

    _replace_unique_sdt(dest_declarante_p, MARCA_SI + ' ')
    _replace_unique_sdt(dest_repr_p, MARCA_NO + ' ')
    _replace_unique_sdt(
        medio_elec_p,
        expr_casilla("notificacion_medio == 'Medios electrónicos'") + ' ',
    )
    _replace_unique_sdt(
        medio_papel_p,
        expr_casilla("notificacion_medio == 'Papel'") + ' ',
    )

    # Sección 4 — Embarcación (tabla 3)
    t3 = tables[3]

    def embarcacion_declarada_expr(label):
        return expr_casilla(f"embarcacion_declarada_antes == '{label}'")

    embarcacion_declarada_opciones = process_checkbox_group(
        t3.rows[0].cells[0].paragraphs[0], embarcacion_declarada_expr
    )
    if len(embarcacion_declarada_opciones) != 2:
        raise RuntimeError(
            'Se esperaban 2 casillas de embarcación declarada/no declarada, '
            f'se encontraron {len(embarcacion_declarada_opciones)}'
        )

    n_cajas = replace_textbox_matching(t3.rows[1].cells[0], 'CHMS', '{{ embarcacion_matricula }}')
    if n_cajas < 1:
        raise RuntimeError('No se encontraron los huecos de matrícula CHMS en el cuadro de texto.')

    append_text(t3.rows[2].cells[2].paragraphs[0], ' {{ embarcacion_tipo }}')
    marca_modelo_cell = t3.rows[2].cells[4]
    append_text(marca_modelo_cell.paragraphs[0], ' {{ embarcacion_marca_modelo }}')
    append_text(marca_modelo_cell.paragraphs[1], ' {{ embarcacion_material_casco }}')
    append_text(marca_modelo_cell.paragraphs[2], ' {{ embarcacion_numero_serie }}')
    append_text(t3.rows[3].cells[0].paragraphs[1], ' {{ embarcacion_eslora }}')
    append_text(t3.rows[3].cells[2].paragraphs[1], ' {{ embarcacion_manga }}')
    append_text(t3.rows[3].cells[4].paragraphs[0], ' {{ embarcacion_ocupantes_max }}')
    append_text(t3.rows[4].cells[1].paragraphs[0], ' {{ embarcacion_vela_altura_mastil }}')
    append_text(t3.rows[5].cells[1].paragraphs[0], ' {{ embarcacion_motor_num_motores }}')
    append_text(t3.rows[5].cells[2].paragraphs[0], ' {{ embarcacion_motor_tipo }}')
    append_text(t3.rows[5].cells[3].paragraphs[0], ' {{ embarcacion_motor_marca }}')
    append_text(t3.rows[5].cells[5].paragraphs[0], ' {{ embarcacion_motor_potencia_cv }}')
    append_text(t3.rows[6].cells[2].paragraphs[0], ' {{ embarcacion_motor_ciclo }}')

    # Sección 5 — Seguro (tabla 4)
    t4 = tables[4]
    set_cell_text(t4.rows[2].cells[0], '{{ seguro_poliza }}')
    replace_paragraph_text(
        t4.rows[2].cells[1].paragraphs[0],
        'De {{ seguro_vigencia_desde }} al {{ seguro_vigencia_hasta }}',
    )

    # Sección 6 — Usuario y actividad (tabla 5)
    t5 = tables[5]

    def usuario_tipo_expr(label):
        return expr_casilla(f"usuario_tipo == '{label}'")

    usuario_tipo_opciones = process_checkbox_group(
        t5.rows[1].cells[0].paragraphs[1], usuario_tipo_expr
    )

    def _actividad(label):
        return label[:-1] if label.endswith(':') else label

    def actividad_tipo_expr(label):
        return expr_casilla(f"actividad_tipo == '{_actividad(label)}'")

    def actividad_sufijo(label):
        return ' {{ actividad_otros }}' if _actividad(label) == 'Otros' else ''

    actividad_tipo_opciones: list[str] = []
    for indice_parrafo in (1, 2):
        actividad_tipo_opciones += process_checkbox_group(
            t5.rows[2].cells[0].paragraphs[indice_parrafo],
            actividad_tipo_expr,
            sufijo_builder=actividad_sufijo,
        )

    # Sección 7A — Embalses (tabla 6), solo lista plazo <= 1 año
    t6 = tables[6]

    def embalse_expr(label):
        return expr_casilla(f'"{label}" in embalses_seleccionados')

    for col in (0, 1, 2, 3):
        cell = t6.rows[3].cells[col]
        for paragraph in cell.paragraphs:
            embalses_taggeados.extend(process_checkbox_group(paragraph, embalse_expr))

    append_text(t6.rows[4].cells[0].paragraphs[0], ' {{ embalses_otros }}')

    # Sección 8 — Plazo opción A (tabla 7)
    t7 = tables[7]
    opcion_a_cell = t7.rows[2].cells[1]
    todo_el_anio_p = opcion_a_cell.paragraphs[1]
    meses_anio_p = opcion_a_cell.paragraphs[2]
    meses_lista_p = opcion_a_cell.paragraphs[3]
    dias_especificos_p = t7.rows[3].cells[1].paragraphs[0]

    if not replace_leading_sym_with_mark(
        todo_el_anio_p, expr_casilla("plazo_modalidad == 'Todo el año'") + ' '
    ):
        raise RuntimeError("No se encontró la casilla de 'TODO EL AÑO'.")
    if not replace_leading_sym_with_mark(
        meses_anio_p, expr_casilla("plazo_modalidad == 'Meses concretos'") + ' '
    ):
        raise RuntimeError("No se encontró la casilla de 'MESES DEL AÑO'.")
    if not replace_leading_sym_with_mark(
        dias_especificos_p, expr_casilla("plazo_modalidad == 'Fechas concretas'") + ' '
    ):
        raise RuntimeError("No se encontró la casilla de 'DÍAS ESPECÍFICOS'.")

    condicionar_anio(todo_el_anio_p, '_ _', "plazo_modalidad == 'Todo el año'")
    condicionar_anio(meses_anio_p, '_ _ ', "plazo_modalidad == 'Meses concretos'")

    def mes_expr(label):
        canon = label.strip().capitalize()
        return expr_casilla(
            f"plazo_modalidad == 'Meses concretos' and '{canon}' in plazo_meses"
        )

    meses_taggeados = process_checkbox_group(meses_lista_p, mes_expr)
    if len(meses_taggeados) != 12:
        raise RuntimeError(
            f'Se esperaban 12 casillas de mes, se encontraron {len(meses_taggeados)}'
        )

    dias_runs = list(dias_especificos_p.runs)
    start_idx = 6
    dias_runs[start_idx].text = (
        "{% if plazo_modalidad == 'Fechas concretas' %}"
        '{{ plazo_fecha_desde }} al {{ plazo_fecha_hasta }}'
        '{% endif %}'
    )
    for extra_run in dias_runs[start_idx + 1 :]:
        extra_run.text = ''

    # Sección 9 — Oposición consulta CHMS (tabla 8)
    t8 = tables[8]
    oposicion_p = t8.rows[1].cells[0].paragraphs[1]
    oposicion_sdts = oposicion_p._p.findall('.//' + qn('w:sdt'))
    if len(oposicion_sdts) != 1:
        raise RuntimeError(
            f'Se esperaba 1 casilla w:sdt para oposición a consulta CHMS, '
            f'se encontraron {len(oposicion_sdts)}'
        )
    replace_sdt_with_text(oposicion_sdts[0], expr_casilla('oposicion_consulta_chms') + ' ')

    # Sección 11 — Firma (tabla anidada en tabla 9)
    t9 = tables[9]
    firma_container_cell = t9.rows[2].cells[0]
    nested_tables = firma_container_cell.tables
    if len(nested_tables) != 1:
        raise RuntimeError(
            f'Se esperaba 1 tabla anidada para el bloque de firma, '
            f'se encontraron {len(nested_tables)}'
        )
    firma_table = nested_tables[0]
    append_text(firma_table.rows[0].cells[1].paragraphs[0], '{{ firma_lugar }}')
    append_text(firma_table.rows[0].cells[3].paragraphs[0], '{{ firma_fecha }}')
    set_cell_text(firma_table.rows[0].cells[4], '')
    set_cell_text(firma_table.rows[0].cells[6], '')
    set_cell_text(firma_table.rows[0].cells[7], '')
    set_cell_text(firma_table.rows[0].cells[8], '')
    append_text(firma_table.rows[2].cells[0].paragraphs[1], ' {{ firma_nombre }}')
    append_text(firma_table.rows[2].cells[5].paragraphs[1], ' {{ firma_dni }}')

    doc.save(str(dst))
    _guardar_metadata(_huella_archivo(src), dst)

    return {
        'plantilla': str(dst),
        'editable': str(src),
        'embalses_taggeados': embalses_taggeados,
        'usuario_tipo_opciones': usuario_tipo_opciones,
        'actividad_tipo_opciones': actividad_tipo_opciones,
        'declarante_tipo_opciones': declarante_tipo_opciones,
        'embarcacion_declarada_opciones': embarcacion_declarada_opciones,
        'meses_taggeados': meses_taggeados,
    }


def _modo_regenerar() -> str:
    valor = os.getenv('REGENERAR_PLANTILLA_ON_START', 'auto').strip().lower()
    if valor in ('never', 'always'):
        return valor
    return 'auto'


def regenerar_plantilla_si_cambio() -> dict:
    """Regenera la plantilla según REGENERAR_PLANTILLA_ON_START. Seguro para lifespan."""
    src = ruta_editable()
    dst = ruta_plantilla()
    modo = _modo_regenerar()

    if modo == 'never':
        msg = 'REGENERAR_PLANTILLA_ON_START=never: no se comprueba ni regenera la plantilla.'
        log.info(msg)
        return {'ok': True, 'action': 'skipped', 'message': msg, 'plantilla': str(dst)}

    if not src.is_file():
        if plantilla_disponible(dst):
            msg = f'No hay DOCX editable en {src}; se conserva {dst}.'
            log.info(msg)
            return {'ok': True, 'action': 'kept_cache', 'message': msg, 'plantilla': str(dst)}
        msg = f'No hay DOCX editable en {src} ni plantilla; la declaración no estará disponible.'
        log.warning(msg)
        return {'ok': False, 'action': 'missing', 'message': msg, 'plantilla': str(dst)}

    if modo == 'auto' and not necesita_regenerar_plantilla(src, dst):
        msg = f'Sin cambios en el formulario editable; {dst} está actualizada.'
        log.info(msg)
        return {'ok': True, 'action': 'up_to_date', 'message': msg, 'plantilla': str(dst)}

    try:
        resultado = build_plantilla(editable=src, plantilla=dst)
    except Exception as exc:
        msg = f'No se pudo regenerar la plantilla DR: {exc}'
        log.exception(msg)
        return {'ok': False, 'action': 'error', 'message': msg, 'plantilla': str(dst)}

    msg = f'Plantilla regenerada en {resultado["plantilla"]}.'
    log.info(msg)
    return {
        'ok': True,
        'action': 'regenerated',
        'message': msg,
        'plantilla': resultado['plantilla'],
        'detalle': resultado,
    }
